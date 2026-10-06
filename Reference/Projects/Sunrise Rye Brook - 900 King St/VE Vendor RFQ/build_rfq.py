"""Sunrise Rye Brook - VE vendor request for pricing (equipment only, no RFI content).

Run: python3 build_rfq.py  ->  VE Vendor RFQ - Sunrise Rye Brook.pdf / .xlsx
Item data comes from ../VE Mark-ups/build_markups.py (VAV decisions, kW, cabinet heaters, unit heaters) and the
GMP VAV schedule (M6000). No prices from us; vendors fill in unit prices in the xlsx.
"""
import os, importlib.util
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
from reportlab.platypus import (BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, Table, TableStyle,
                                KeepTogether)
from reportlab.lib.styles import ParagraphStyle
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

S = os.path.dirname(os.path.abspath(__file__)) + "/"
spec = importlib.util.spec_from_file_location("mk", S + "../VE Mark-ups/build_markups.py")
mk = importlib.util.module_from_spec(spec); spec.loader.exec_module(mk)   # registers Arial fonts too

OUT_PDF = S + "VE Vendor RFQ - Sunrise Rye Brook.pdf"
OUT_XLS = S + "VE Vendor RFQ - Sunrise Rye Brook.xlsx"
DATE, REV = "10/6/2026", "Rev 0"
PW, PH = letter
BLUE = colors.HexColor("#236fa1"); GREY = colors.HexColor("#aeaeae"); LIGHT = colors.HexColor("#eef3f8")
LINE = colors.HexColor("#c7ced6")

# GMP VAV schedule M6000: tag -> (inlet size, max CFM, min CFM)
VSCH = {1: (12, 1200, 825), 2: (14, 1450, 1350), 3: (14, 2250, 1150), 4: (8, 700, 375), 5: (8, 400, 400),
        6: (8, 400, 400), 7: (10, 700, 250), 8: (4, 200, 100), 9: (8, 350, 150), 10: (6, 425, 250),
        11: (10, 550, 500), 12: (14, 1300, 1300), 13: (10, 600, 600), 14: (10, 600, 600), 15: (6, 200, 200),
        16: (6, 300, 250), 17: (8, 450, 450), 18: (8, 500, 500), 19: (6, 200, 50), 20: (8, 350, 150),
        21: (8, 300, 150), 22: (12, 900, 250), 23: (6, 300, 100), 24: (10, 550, 200), 25: (14, 3000, 1100),
        26: (6, 400, 200), 27: (14, 1850, 425), 28: (10, 800, 525), 29: (16, 1800, 600), 30: (6, 350, 200),
        31: (14, 1200, 400), 32: (6, 150, 50), 33: (16, 2500, 1000), 34: (14, 1200, 450), 35: (6, 200, 200),
        36: (8, 350, 150), 37: (8, 250, 150), 38: (8, 450, 300), 39: (6, 400, 250), 40: (6, 200, 100),
        41: (8, 400, 150), 42: (12, 600, 200), 43: (14, 1000, 400), 44: (16, 1100, 500), 45: (8, 450, 350),
        46: (6, 400, 200), 47: (6, 450, 300)}
VAV10_KW = 3.0   # VAV-10 (locker area) keeps a box; with the boiler plant deleted it needs electric reheat too

# ---------------------------------------------------------------- item lists
vav_e, vav_c, vav_del = [], [], []
for r in mk.VAV:
    n = int(r[1].split("-")[1])
    size, mx, mn = VSCH[n]
    if r[2] == "EREHEAT":
        if r[3].startswith("Treatment"):
            vav_e.append(["VAV-18A", r[3], "TBD", "TBD", "TBD", mk.vav_kw(r)])
        else:
            vav_e.append([r[1], r[3], size, mx, mn, mk.vav_kw(r)])
    elif r[2] == "NOREHEAT":
        vav_c.append([r[1], r[3], size, mx, mn])
    elif r[2] == "CANCEL":
        vav_del.append([r[1], r[3], size, mx, mn])
    elif r[2] == "NC":
        vav_e.append([r[1], "Locker rooms", size, mx, mn, VAV10_KW])
ch = [[t, room, kw] for p, t, room, kw in mk.CH]
euh = [[room, kw] for p, t, x, y, act, room, kw in mk.FCU if act == "EUH"]

DELETED = [
  ("Boilers B-1, B-2, B-3 (1,999 MBH input, fire tube)", 3, "ea"),
  ("Boiler pumps BP-1, BP-2", 2, "ea"),
  ("Primary pumps PP-1, PP-2", 2, "ea"),
  ("Glycol / snow melt pumps P-1 to P-4", 4, "ea"),
  ("Heat exchanger HX-3 (826 MBH, 30% PG)", 1, "ea"),
  ("Hydraulic separator HS-1, expansion tanks ET-1 / ET-2, air separator AS-1", 1, "lot"),
  ("Chemical pot feeder and glycol fill package", 1, "lot"),
  ("Hot water fan coil units FCU-A (basement)", 20, "ea"),
  ("Radiators RAD-1, RAD-2", 2, "ea"),
  ("Hot water cabinet heaters CH-1, CH-3 to CH-12", 11, "ea"),
  ("Hot water reheat coils on VAV boxes", 47, "ea"),
  ("VAV boxes cancelled (list in section 1C)", len(vav_del), "ea"),
]
ALTERNATES = [
  ("ALT-1", "Deduct: combination fire / smoke dampers at apartment DOAS branches, 6x4 to 8x6, with actuators "
            "(floors 1 to 4)", 283, "ea"),
  ("ALT-2", "Deduct: apartment bath exhaust fans TX-A (75 CFM, ceiling, 120 V)", 283, "ea"),
]

# ---------------------------------------------------------------- PDF
pB = ParagraphStyle("b", fontName="Arial", fontSize=9.5, leading=12)
pS = ParagraphStyle("s", fontName="Arial", fontSize=8, leading=9.8)
pH = ParagraphStyle("h", fontName="Arial-Bold", fontSize=10.5, leading=13, textColor=BLUE, spaceBefore=8, spaceAfter=4)
pH2 = ParagraphStyle("h2", fontName="Arial-Bold", fontSize=9, leading=11, spaceBefore=6, spaceAfter=3)


def first(c, doc):
    img = ImageReader(mk.LOGO); w = 113; h = w * 1303 / 2367
    c.drawImage(img, 45, PH - 32 - h, width=w, height=h, mask="auto")
    c.setFont("Arial", 9); c.drawRightString(PW - 45, PH - 40, "Request for Pricing / Date: %s / %s" % (DATE, REV))
    y = PH - 118
    for line in ["Northern Wolves AC", "55 9th St, 55-A2", "Brooklyn, NY", "11215, US", "(347) 463-9248", "",
                 "Requested By:", "Alikhan Ilyas", "(347) 463-9248", "alikhan@northernwolvesac.com"]:
        c.drawString(45, y, line); y -= 10.5
    c.setFont("Arial", 8); c.setFillColor(GREY); c.drawRightString(PW - 45, PH - 118, "To")
    c.setFont("Arial", 9); c.setFillColor(colors.black); c.drawRightString(PW - 45, PH - 128, "Equipment vendors")
    foot(c, doc)


def later(c, doc):
    c.setFont("Arial", 8); c.setFillColor(GREY)
    c.drawString(45, PH - 30, "Sunrise Rye Brook - VE request for pricing %s - %s" % (REV, DATE))
    c.setFillColor(colors.black); foot(c, doc)


def foot(c, doc):
    c.setFont("Arial", 7.5); c.setFillColor(GREY)
    c.drawString(45, 25, "Northern Wolves AC - vendor request for pricing. Please quote per the attached xlsx.")
    c.drawRightString(PW - 45, 25, "Page %d" % doc.page); c.setFillColor(colors.black)


def tbl(rows, widths, total=False):
    t = Table(rows, colWidths=widths, repeatRows=1)
    st = [("FONT", (0, 0), (-1, 0), "Arial-Bold", 7.8), ("FONT", (0, 1), (-1, -1), "Arial", 7.8),
          ("BACKGROUND", (0, 0), (-1, 0), LIGHT), ("VALIGN", (0, 0), (-1, -1), "TOP"),
          ("GRID", (0, 0), (-1, -1), 0.4, LINE), ("TOPPADDING", (0, 0), (-1, -1), 2),
          ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("ALIGN", (2, 1), (-1, -1), "RIGHT")]
    if total:
        st += [("FONT", (0, -1), (-1, -1), "Arial-Bold", 7.8), ("BACKGROUND", (0, -1), (-1, -1), LIGHT)]
    t.setStyle(TableStyle(st)); return t


doc = BaseDocTemplate(OUT_PDF, pagesize=letter, leftMargin=45, rightMargin=45, topMargin=45, bottomMargin=40,
                      title="Sunrise Rye Brook - VE Request for Pricing", author="Northern Wolves AC")
doc.addPageTemplates([PageTemplate("first", [Frame(45, 40, PW - 90, PH - 278)], onPage=first, autoNextPageTemplate="later"),
                      PageTemplate("later", [Frame(45, 40, PW - 90, PH - 85)], onPage=later)])
s = [Paragraph("Project: <b>Sunrise Rye Brook - 900 King St, Rye Brook, NY 10573</b>", pB), Spacer(1, 3),
     Paragraph("<b>Mechanical drawings: GMP set dated 07.15.2026 (value engineering revisions below)</b>", pB),
     Paragraph("Request for revised pricing - value engineering", pH),
     Paragraph("The heating design is being revised from a hot water boiler plant to electric heat. Please provide "
               "revised pricing for the items below: new items (section 1 and 2), credits for items deleted from your "
               "previous quote (section 3) and alternate deducts (section 4). Quote per line in the attached xlsx, "
               "with lead times. Electric heat voltage to be confirmed by the electrical engineer: quote 277 V / 1 ph "
               "up to 5 kW and 480 V / 3 ph above 5 kW, and note any other standard voltage. Controls are standalone "
               "(no BMS): factory-mounted standalone controllers with wall thermostats / sensors.", pS)]

s += [Paragraph("1. VAV terminal units (1st floor common areas)", pH),
      Paragraph("1A. Single-duct VAV with factory electric reheat coil, SCR control, airflow switch, disconnect, "
                "standalone pressure-independent controller and wall thermostat (%d)" % len(vav_e), pH2)]
rows = [["Tag", "Room served", "Inlet", "Max CFM", "Min CFM", "Elec. heat"]]  # noqa
for t, room, sz, mx, mn, kw in vav_e:
    rows.append([t, Paragraph(room, pS), ('%s"' % sz) if sz != "TBD" else "TBD", mx, mn, "%g kW" % kw])
rows.append(["Total", "", "", "", "", "%g kW" % sum(r[5] for r in vav_e)])
s += [tbl(rows, [50, 220, 45, 60, 60, 87], total=True),
      Paragraph("Heating airflow at the scheduled minimum CFM; confirm minimum airflow for the electric coil. "
                "VAV-18A: second VAV-18 tag on M2010 (Treatment Rm 1041C), size per vendor at 3 kW.", pS)]
s += [Paragraph("1B. Single-duct VAV, cooling only (no reheat coil), standalone controller and wall thermostat (%d)"
                % len(vav_c), pH2)]
rows = [["Tag", "Room served", "Inlet", "Max CFM", "Min CFM"]] + \
       [[t, Paragraph(room, pS), '%s"' % sz, mx, mn] for t, room, sz, mx, mn in vav_c]
s += [tbl(rows, [50, 280, 45, 70, 77])]
s += [Paragraph("1C. VAV boxes cancelled - credit (%d)" % len(vav_del), pH2)]
rows = [["Tag", "Room served", "Inlet", "Max CFM", "Min CFM"]] + \
       [[t, Paragraph(room, pS), '%s"' % sz, mx, mn] for t, room, sz, mx, mn in vav_del]
s += [tbl(rows, [50, 280, 45, 70, 77])]

s += [Paragraph("2. Electric heaters", pH),
      Paragraph("2A. Electric cabinet unit heaters, wall / ceiling as scheduled, with integral thermostat, "
                "disconnect and fan (%d)" % len(ch), pH2)]
rows = [["Tag", "Location", "Heat", "Fan CFM (as scheduled)"]] + \
       [[t, room, "%g kW" % kw, "1,050" if t == "CH-9" else "450"] for t, room, kw in ch]
rows.append(["Total", "", "%g kW" % sum(c[2] for c in ch), ""])
s += [tbl(rows, [50, 250, 80, 142], total=True)]
s += [Paragraph("2B. Electric horizontal unit heaters, ceiling hung, with integral thermostat and disconnect (%d)"
                % len(euh), pH2)]
rows = [["#", "Location (basement)", "Heat"]] + [[str(i), room, "%g kW" % kw] for i, (room, kw) in enumerate(euh, 1)]
rows.append(["Total", "", "%g kW" % sum(e[1] for e in euh)])
s += [tbl(rows, [50, 330, 142], total=True)]

s += [Paragraph("3. Deleted from scope - please provide credit against your previous quote", pH)]
rows = [["Item", "Qty", "Unit"]] + [[Paragraph(d, pS), q, u] for d, q, u in DELETED]
s += [tbl(rows, [400, 60, 62])]
s += [Paragraph("4. Alternate deducts - price separately", pH)]
rows = [["Alt", "Item", "Qty", "Unit"]] + [[a, Paragraph(d, pS), q, u] for a, d, q, u in ALTERNATES]
s += [tbl(rows, [40, 360, 60, 62])]
s += [Spacer(1, 6), Paragraph("Please include: unit and extended prices, freight to Rye Brook NY, lead time, "
                              "submittal data, and any exclusions. Taxes listed separately.", pS)]
s += [Spacer(1, 6), Paragraph("Attached: marked-up drawings M2010, M2011, M2013 (VAV terminals, cabinet heaters) and "
                              "M2000, M2002 (basement unit heaters). Each item is marked with its tag and heat kW.", pS)]
doc.build(s)

# append the marked-up sheets (vendor copy, no RFI references)
import pymupdf
rfq = pymupdf.open(OUT_PDF)
rfq.insert_pdf(mk.build_sheets(vendor=True))
rfq.set_metadata({"title": "Sunrise Rye Brook - VE Request for Pricing", "author": "Northern Wolves AC"})
rfq.save(OUT_PDF + ".tmp", garbage=3, deflate=True); os.replace(OUT_PDF + ".tmp", OUT_PDF)

# ---------------------------------------------------------------- XLSX
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Pricing"
B = Font(bold=True); H = PatternFill("solid", fgColor="EEF3F8"); thin = Side(style="thin", color="C7CED6")
ws.append(["Sunrise Rye Brook - VE request for pricing %s %s - Northern Wolves AC" % (REV, DATE)]); ws["A1"].font = Font(bold=True, size=12)
ws.append(["Fill in unit price and lead time. Electric heat: 277 V/1 ph up to 5 kW, 480 V/3 ph above (to be confirmed)."])
ws.append([])
hdr = ["Section", "Tag / item", "Description", "Inlet", "Max CFM", "Min CFM", "Elec. heat kW", "Qty", "Unit",
       "Unit price", "Extended", "Lead time", "Notes"]


def section(title, items):
    ws.append([title]); ws.cell(ws.max_row, 1).font = B
    ws.append(hdr)
    for c in ws[ws.max_row]: c.font = B; c.fill = H
    for it in items:
        ws.append(it)
        r = ws.max_row
        ws.cell(r, 11).value = "=IF(J%d=\"\",\"\",H%d*J%d)" % (r, r, r)
    ws.append([])


section("1A VAV with electric reheat", [["1A", t, room, sz, mx, mn, kw, 1, "ea"] for t, room, sz, mx, mn, kw in vav_e])
section("1B VAV cooling only", [["1B", t, room, sz, mx, mn, "", 1, "ea"] for t, room, sz, mx, mn in vav_c])
section("1C VAV cancelled - credit", [["1C", t, room, sz, mx, mn, "", -1, "ea"] for t, room, sz, mx, mn in vav_del])
section("2A Electric cabinet heaters", [["2A", t, room, "", "1,050" if t == "CH-9" else "450", "", kw, 1, "ea"] for t, room, kw in ch])
section("2B Electric unit heaters", [["2B", "EUH-%d" % i, room, "", "", "", kw, 1, "ea"] for i, (room, kw) in enumerate(euh, 1)])
section("3 Deleted - credit", [["3", "", d, "", "", "", "", -q, u] for d, q, u in DELETED])
section("4 Alternate deducts", [["4", a, d, "", "", "", "", -q, u] for a, d, q, u in ALTERNATES])
for col, w in zip("ABCDEFGHIJKLM", [9, 11, 52, 7, 9, 9, 12, 6, 6, 12, 13, 12, 30]):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=4):
    for c in row:
        if c.value is not None and row[0].value in ("1A", "1B", "1C", "2A", "2B", "3", "4", "Section"):
            c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            c.alignment = Alignment(vertical="top", wrap_text=(c.column == 3))
ws.freeze_panes = "A4"
wb.save(OUT_XLS)
print(OUT_PDF); print(OUT_XLS)
print("1A", len(vav_e), sum(r[5] for r in vav_e), "kW; 1B", len(vav_c), "; 1C", len(vav_del), "; CH", len(ch), "; EUH", len(euh))
