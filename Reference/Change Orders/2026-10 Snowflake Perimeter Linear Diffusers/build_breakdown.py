"""Snowflake 7 Times Square fl. 29-31: perimeter linear diffuser revision - quantity breakdown per floor (no pricing).

Source: our shop drawings M-529 / M-530 / M-531 SD R2 (sheets A = north half, B = south half), revision clouds.
Each bubbled area: perimeter linear diffuser(s) in plenum box O.D. 48x4x26 with 1/2" lining, 12x10 branch with VD/CO
off the perimeter main. Cloud rectangles are found from the red cloud arcs in the PDFs (see find_clouds()).

Run:  python3 build_breakdown.py
Out:  Perimeter Linear Diffusers - Breakdown by Floor.pdf / .xlsx next to this script.
"""
import io, os
import pymupdf
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, TABLOID
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, Table, TableStyle
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

pymupdf.TOOLS.mupdf_display_errors(False)
HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'source')
SHEETS = {29: 'M-529 SD R2 (10-02-2026).pdf', 30: 'M-530 SD R2 (10-02-2026).pdf', 31: 'M-531 SD R2.pdf'}
DOCS = {f: pymupdf.open(os.path.join(SRC, n)) for f, n in SHEETS.items()}
OUT = os.path.join(HERE, 'Perimeter Linear Diffusers - Breakdown by Floor')
DATE = '10/05/2026'
PLAN = (540, 150, 3250, 2120)           # plan area on the SD sheets (rotated page coords)

# cloud = index of the cloud on that sheet in find_clouds() order (north to south, west to east)
# cfm = one entry per linear diffuser; br = (branch size, length in inches) per connection
# main = mains inside the bubbled area (size, approx LF), for reference - see open questions
AREAS = [
    dict(id='29-1', floor=29, sheet='A', cloud=0, loc='North-east corner, W 42nd St / Broadway',
         cfm=[300, 300, 300, 300], br=[('12x10', 4), ('12x10', 7), ('12x10', 7), ('12x10', 7)], vdco=4,
         main='31.3/12 approx. 16 LF; 26.1/10 with offset down approx. 6 LF; 14"Ø approx. 6 LF',
         note='2 diffusers along 42nd St off the 31.3/12 main, 2 along Broadway off the 26.1/10 / 14"Ø run.'),
    dict(id='29-2', floor=29, sheet='A', cloud=1, loc='Seventh Ave, north',
         cfm=[265, 270, 265], br=[('12x10', 11)] * 3, vdco=3,
         main='29.3/10 approx. 13 LF', note=''),
    dict(id='29-3', floor=29, sheet='A', cloud=2, loc='Seventh Ave, Open Office 29.42',
         cfm=[265, 265, 270, 265, 270, 265], br=[('12x10', 11)] * 6, vdco=2,
         main='26.1/10 approx. 20 LF (with tee)', note='VD/CO tagged on 2 of 6 connections.'),
    dict(id='29-4', floor=29, sheet='A', cloud=3, loc='Broadway, north',
         cfm=[200, 200], br=[('12x10', 7)] * 2, vdco=2,
         main='34.4/12 approx. 8 LF', note=''),
    dict(id='29-5', floor=29, sheet='B', cloud=0, loc='Seventh Ave, south (north end of sheet B)',
         cfm=[300, 300, 300, 300], br=[('12x10', 10)] * 4, vdco=4,
         main='31.3/12 approx. 12 LF; 14"Ø approx. 8 LF', note=''),
    dict(id='29-6', floor=29, sheet='B', cloud=1, loc='Seventh Ave, south',
         cfm=[300, 300, 300, 300], br=[('12x10', 10), ('12x10', 10), ('12x10', 10), ('10x8', 10)], vdco=3,
         main='14"Ø approx. 13 LF; 29.3/10 approx. 7 LF', note='Last connection 10x8. VD/CO tagged on 3 of 4.'),

    dict(id='30-1', floor=30, sheet='A', cloud=0, loc='Seventh Ave, north',
         cfm=[265, 270, 265], br=[('12x10', 11)] * 3, vdco=3,
         main='26.1/10 approx. 12 LF', note=''),
    dict(id='30-2', floor=30, sheet='A', cloud=1, loc='Seventh Ave, Open Office 30.35',
         cfm=[300, 300, 300, 200, 200], br=[('12x10', 11)] * 3 + [('12x10', 7)] * 2, vdco=5,
         main='26.1/10 approx. 12 LF; 12"Ø approx. 9 LF', note='3 off the 26.1/10, 2 off the 12"Ø below the 19x15.'),
    dict(id='30-3', floor=30, sheet='B', cloud=0, loc='South wall (W 41st St side)',
         cfm=[300, 300, 300, 300], br=[('12x10', 11)] * 4, vdco=4,
         main='32.2/16 approx. 40 LF', note=''),

    dict(id='31-1', floor=31, sheet='A', cloud=0, loc='Seventh Ave, north (Focus Rm 31.65 / Corridor 31.51)',
         cfm=[200, 200], br=[('12x10', 9)] * 2, vdco=1,
         main='26.1/10 approx. 4 LF (at elbow)', note=''),
    dict(id='31-2', floor=31, sheet='A', cloud=1, loc='Seventh Ave, Huddle Rms 31.46 / 31.47',
         cfm=[265, 270, 265], br=[('12x10', 12)] * 3, vdco=0,
         main='28.1/12 approx. 15 LF', note='No VD/CO tags shown.'),
    dict(id='31-3', floor=31, sheet='A', cloud=2, loc='Seventh Ave, at IDF 31.44 (match line)',
         cfm=[300], br=[('12x10', 9)], vdco=0,
         main='34.4/12 (same run as 31-4)', note='No VD/CO tag shown.'),
    dict(id='31-4', floor=31, sheet='B', cloud=0, loc='Seventh Ave, 7P Meeting Rm 31.43',
         cfm=[300, 300, 300], br=[('12x10', 9)] * 3, vdco=0,
         main='34.4/12 approx. 20 LF', note='No VD/CO tags shown.'),
    dict(id='31-5', floor=31, sheet='B', cloud=1, loc='Seventh Ave, south',
         cfm=[265, 270, 265], br=[('12x10', 9)] * 3, vdco=0,
         main='28.1/12 approx. 15 LF', note='No VD/CO tags shown.'),
    dict(id='31-6', floor=31, sheet='B', cloud=2, loc='Broadway, south',
         cfm=[300, 300], br=[('12x10', 7)] * 2, vdco=2,
         main='29.3/10 approx. 13 LF', note=''),
]
FLOORS = [29, 30, 31]
# check against the counts read from the drawing text (PLENUM labels / CFM tags in the clouds)
assert [len(a['cfm']) for a in AREAS] == [4, 3, 6, 2, 4, 4, 3, 5, 4, 2, 3, 1, 3, 3, 2]
assert all(len(a['br']) == len(a['cfm']) for a in AREAS)


def find_clouds(page, tol=12):
    """Revision clouds = red paths made only of curve segments; merge touching ones, drop small symbols."""
    rs = [pymupdf.Rect(d['rect']) for d in page.get_drawings()
          if d.get('color') and tuple(round(x, 2) for x in d['color']) == (1.0, 0.0, 0.0)
          and set(i[0] for i in d['items']) == {'c'}]
    cl = []
    for r in rs:
        r = pymupdf.Rect(r.x0 - tol, r.y0 - tol, r.x1 + tol, r.y1 + tol)
        merged = True
        while merged:
            merged = False
            for c in cl:
                if c.intersects(r):
                    r |= c; cl.remove(c); merged = True; break
        cl.append(r)
    out = []
    for c in cl:
        c = pymupdf.Rect(c.x0 + tol, c.y0 + tol, c.x1 - tol, c.y1 - tol)
        if c.width > 100 and c.height > 100:
            out.append(pymupdf.Rect(c * page.rotation_matrix).normalize())
    return sorted(out, key=lambda r: (round(r.y0 / 50), r.x0))


CLOUDS = {(f, s): find_clouds(DOCS[f]['AB'.index(s)]) for f in FLOORS for s in 'AB'}
assert {k: len(v) for k, v in CLOUDS.items()} == {(29, 'A'): 4, (29, 'B'): 2, (30, 'A'): 2, (30, 'B'): 1, (31, 'A'): 3, (31, 'B'): 3}, \
    {k: len(v) for k, v in CLOUDS.items()}
assert all(sum(1 for a in AREAS if (a['floor'], a['sheet']) == k) == len(v) for k, v in CLOUDS.items())
for a in AREAS:
    a['rect'] = CLOUDS[(a['floor'], a['sheet'])][a['cloud']]
    a['n'] = len(a['cfm'])


def floor_areas(f):
    return [a for a in AREAS if a['floor'] == f]


def brsum(areas):
    out = {}
    for a in areas:
        for size, ln in a['br']:
            n, l = out.get(size, (0, 0)); out[size] = (n + 1, l + ln)
    return out


# ---------------- page layout ----------------
pdfmetrics.registerFont(TTFont('LS', '/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf'))
pdfmetrics.registerFont(TTFont('LSB', '/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf'))
pdfmetrics.registerFontFamily('LS', normal='LS', bold='LSB', italic='LS', boldItalic='LSB')
PW, PH = landscape(TABLOID)
M = 30
NAVY = colors.HexColor('#1F3A5F'); ORANGE = colors.HexColor('#E08A00'); GRAY = colors.HexColor('#555555')
LIGHT = colors.HexColor('#EEF2F7')
st = ParagraphStyle('b', fontName='LS', fontSize=9, leading=11.5)
str_ = ParagraphStyle('r', parent=st, alignment=2)
sth = ParagraphStyle('h', fontName='LSB', fontSize=10.5, leading=13, textColor=NAVY)
ordn = lambda n: f'{n}' + ('st' if n % 10 == 1 and n % 100 != 11 else 'nd' if n % 10 == 2 and n % 100 != 12 else 'rd' if n % 10 == 3 and n % 100 != 13 else 'th')
esc = lambda t: t.replace('&', '&amp;')
TOP, BOT = 62, 34
NPAGES = 1 + 2 * len(FLOORS)
page_no = [0]


def header(c, title, sub):
    c.setFillColor(NAVY); c.rect(0, PH - 52, PW, 52, stroke=0, fill=1)
    c.setFillColor(colors.white)
    c.setFont('LSB', 16); c.drawString(M, PH - 26, title)
    c.setFont('LS', 10); c.drawString(M, PH - 42, sub)
    c.setFont('LSB', 11); c.drawRightString(PW - M, PH - 26, 'NORTHERN WOLVES AC')
    c.setFont('LS', 9); c.drawRightString(PW - M, PH - 42, 'Snowflake Fitout - 7 Times Square, Floors 29-31  |  Perimeter Linear Diffusers')


def footer(c):
    page_no[0] += 1
    c.setFillColor(colors.gray); c.setFont('LS', 8)
    c.drawString(M, 16, f'Northern Wolves AC  |  55 9th St, 55-A2, Brooklyn, NY 11215  |  (347) 463-9248  |  {DATE}  |  '
                        'Quantities taken from shop drawings M-529 / M-530 / M-531 SD R2 - no pricing')
    c.drawRightString(PW - M, 16, f'Page {page_no[0]} of {NPAGES}')


def table(c, data, x, ytop, widths, bold_last=True, extra=()):
    t = Table(data, colWidths=widths)
    s = [('FONT', (0, 0), (-1, -1), 'LS', 9), ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
         ('VALIGN', (0, 0), (-1, -1), 'TOP'), ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#9AA5B1')),
         ('TOPPADDING', (0, 0), (-1, -1), 3), ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
         ('BACKGROUND', (0, 0), (-1, 0), NAVY), ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
         ('FONT', (0, 0), (-1, 0), 'LSB', 8.5), ('ALIGN', (0, 0), (-1, 0), 'CENTER'), ('VALIGN', (0, 0), (-1, 0), 'MIDDLE')]
    if bold_last:
        s += [('FONT', (0, -1), (-1, -1), 'LSB', 9), ('BACKGROUND', (0, -1), (-1, -1), LIGHT)]
    t.setStyle(TableStyle(s + list(extra)))
    w, h = t.wrapOn(c, 0, 0); t.drawOn(c, x, ytop - h)
    return h


def para(c, text, style, x, ytop, w):
    p = Paragraph(text, style); _, h = p.wrap(w, 1000); p.drawOn(c, x, ytop - h); return h


def fit(rect, box):
    r = pymupdf.Rect(rect); bx = pymupdf.Rect(box)
    s = min(bx.width / r.width, bx.height / r.height)
    w, h = r.width * s, r.height * s
    x0 = bx.x0 + (bx.width - w) / 2; y0 = bx.y0 + (bx.height - h) / 2
    return pymupdf.Rect(x0, y0, x0 + w, y0 + h)


def png(page, clip, dpi):
    pix = page.get_pixmap(dpi=dpi, clip=clip)
    buf = io.BytesIO(pix.tobytes('png')); buf.seek(0)
    return buf


def crop_rect(a, pad=45):
    return pymupdf.Rect(a['rect']) + (-pad, -pad, pad, pad)


# ---- cover / summary ----
def summary_page(c):
    header(c, 'Perimeter Linear Diffusers - Breakdown by Floor',
           'Revision clouds on shop drawings M-529 / M-530 / M-531 SD R2  |  Quantities only, pricing to follow')
    x = M; y = PH - TOP
    y -= para(c, 'Scope of the revision', sth, x, y, 560) + 4
    notes = [
        'Perimeter linear supply diffusers along the window walls, each in a <b>plenum box O.D. 48x4x26 with 1/2" lining</b> '
        '(4 ft), fed by a short <b>12x10 branch</b> (one 10x8) with <b>VD / CO</b> off the perimeter supply main.',
        'Per the sketching note on the shop drawings: coordinate soffit openings with the GC; field-adjust branch duct '
        'offsets as needed to clear existing framing and align with linear diffuser locations.',
        'Each bubbled area is numbered by floor and shown with its shop drawing screenshot, followed by a quantity page '
        'per floor. Branch lengths are as dimensioned on the shop drawings; main duct lengths inside the clouds are '
        'approximate and listed for reference only.',
        'This breakdown carries no pricing; pricing follows on the basis below.',
    ]
    for i, n in enumerate(notes, 1):
        y -= para(c, f'{i}. {n}', st, x, y, 560) + 5
    y -= 8
    y -= para(c, 'Basis for pricing', sth, x, y, 560) + 4
    qs = [
        '<b>Confirmed 10/05:</b> the Bulletin 3 ceiling linear diffusers (G) at these locations are deleted. No credit: '
        'all original G diffusers have been delivered and the factory will not take them back.',
        '<b>Confirmed 10/05:</b> perimeter mains stay as installed. Added work starts at the new taps on the mains.',
        '<b>Confirmed 10/05:</b> linear diffusers and plenum boxes furnished and installed by us. VD/CO on every '
        'connection (49). CO number not assigned yet.',
        'Soffit openings, access and patching by GC. M-531 title block still dated 08/17/2026.',
    ]
    for i, n in enumerate(qs, 1):
        y -= para(c, f'{i}. {n}', st, x, y, 560) + 5

    x = 640; y = PH - TOP; w = PW - M - x
    y -= para(c, 'Summary by floor', sth, x, y, w) + 4
    rows = [['Floor', 'Bubbled\nareas', 'Linear diffusers /\nplenum boxes', 'Supply\nCFM', '12x10\nbranches',
             '10x8\nbranches', 'VD / CO\n(tagged)']]
    tot = [0] * 6
    for f in FLOORS:
        ar = floor_areas(f); b = brsum(ar)
        v = [len(ar), sum(a['n'] for a in ar), sum(sum(a['cfm']) for a in ar), b.get('12x10', (0, 0))[0],
             b.get('10x8', (0, 0))[0], sum(a['vdco'] for a in ar)]
        tot = [t + x_ for t, x_ in zip(tot, v)]
        rows.append([f'{ordn(f)} floor (M-5{f})'] + [f'{x_:,}' if x_ else '-' for x_ in v])
    rows.append(['Total'] + [f'{x_:,}' for x_ in tot])
    y -= table(c, rows, x, y, [118, 56, 92, 60, 60, 60, 66]) + 16
    y -= para(c, 'Bubbled areas', sth, x, y, w) + 4
    rows = [['Area', 'Location', 'Diffusers', 'CFM']]
    for a in AREAS:
        rows.append([a['id'], Paragraph(esc(a['loc']), st), a['n'], f'{sum(a["cfm"]):,}'])
    rows.append(['', 'Total', sum(a['n'] for a in AREAS), f'{sum(sum(a["cfm"]) for a in AREAS):,}'])
    table(c, rows, x, y, [44, 330, 72, 66])
    footer(c); c.showPage()


# ---- page 1 of each floor: screenshots of the bubbled areas ----
G = 10; TB = 16


def layout(areas):
    """Split the areas, in order, into rows; pick the split that gives the largest common drawing scale.
    Returns [(area, x, y, w, h)] cells in top-left page coords."""
    import itertools
    W, H = PW - 2 * M, PH - TOP - BOT
    n = len(areas); best = None
    for cuts in itertools.product([0, 1], repeat=n - 1):
        rows, cur = [], [areas[0]]
        for a, cut in zip(areas[1:], cuts):
            if cut: rows.append(cur); cur = [a]
            else: cur.append(a)
        rows.append(cur)
        over = len(rows) * (TB + 8) + (len(rows) - 1) * G
        s = (H - over) / sum(max(crop_rect(a).height for a in r) for r in rows)
        for r in rows:
            s = min(s, (W - (len(r) - 1) * G - 6 * len(r)) / sum(crop_rect(a).width for a in r))
        if best is None or s > best[0] + 1e-9:
            best = (s, rows)
    s, rows = best
    hs = [s * max(crop_rect(a).height for a in r) + TB + 8 for r in rows]
    extra = (H - (len(rows) - 1) * G - sum(hs)) / len(rows)
    cells, y = [], TOP
    for r, h in zip(rows, hs):
        h += extra
        ws = [s * crop_rect(a).width + 6 for a in r]
        k = (W - (len(r) - 1) * G) / sum(ws)
        x = M
        for a, w in zip(r, ws):
            cells.append((a, x, y, w * k, h)); x += w * k + G
        y += h + G
    return cells


def screens_page(c, f):
    areas = floor_areas(f)
    header(c, f'{ordn(f)} Floor - M-5{f} SD R2: Bubbled Areas',
           f'{sum(a["n"] for a in areas)} perimeter linear diffusers in plenum boxes O.D. 48x4x26 w/ 1/2" lining, '
           f'{sum(sum(a["cfm"]) for a in areas):,} CFM  |  quantities on next page')
    for a, cx, cy, cw, ch in layout(areas):
        c.setStrokeColor(GRAY); c.setLineWidth(0.8); c.rect(cx, PH - cy - ch, cw, ch)
        c.setFillColor(ORANGE); c.rect(cx, PH - cy - TB, cw, TB, stroke=0, fill=1)
        c.setFillColor(colors.white); c.setFont('LSB', 9)
        c.drawString(cx + 5, PH - cy - 11.5, f'{a["id"]}  {a["loc"]}' if cw > 330 else a['id'])
        c.drawRightString(cx + cw - 5, PH - cy - 11.5, f'{a["n"]} x LD  |  {sum(a["cfm"]):,} CFM')
        cr = crop_rect(a); p = DOCS[f]['AB'.index(a['sheet'])]
        box = pymupdf.Rect(cx + 3, cy + TB + 4, cx + cw - 3, cy + ch - 4)
        r = fit(cr, box)
        dpi = int(min(300, max(110, 72 * r.width / cr.width * 2.2)))
        c.drawImage(ImageReader(png(p, cr, dpi)), r.x0, PH - r.y1, r.width, r.height)
        c.setFont('LS', 7.5); c.setFillColor(GRAY)
        c.drawRightString(cx + cw - 5, PH - cy - ch + 4, f'M-5{f}{a["sheet"]}.SD')
    footer(c); c.showPage()


# ---- page 2 of each floor: quantities ----
def key_plans(c, f, x, ytop, w):
    h_tot = 0
    for s in 'AB':
        p = DOCS[f]['AB'.index(s)]
        r = fit(PLAN, pymupdf.Rect(x, PH - ytop + h_tot, x + w, PH - ytop + h_tot + 260))
        c.drawImage(ImageReader(png(p, pymupdf.Rect(PLAN), 40)), r.x0, PH - r.y1, r.width, r.height)
        c.setStrokeColor(colors.black); c.setLineWidth(0.5); c.rect(r.x0, PH - r.y1, r.width, r.height)
        c.setFillColor(NAVY); c.setFont('LSB', 8); c.drawString(r.x0 + 3, PH - r.y1 + 4, f'M-5{f}{s}.SD')
        sc = r.width / (PLAN[2] - PLAN[0])
        for a in [a for a in floor_areas(f) if a['sheet'] == s]:
            q = a['rect']
            ax0 = r.x0 + (q.x0 - PLAN[0]) * sc; ay0 = r.y0 + (q.y0 - PLAN[1]) * sc
            ax1 = r.x0 + (q.x1 - PLAN[0]) * sc; ay1 = r.y0 + (q.y1 - PLAN[1]) * sc
            c.setStrokeColor(ORANGE); c.setLineWidth(1.4); c.rect(ax0, PH - ay1, ax1 - ax0, ay1 - ay0)
            c.setFillColor(ORANGE); c.rect(ax0, PH - ay0, 24, 10, stroke=0, fill=1)
            c.setFillColor(colors.white); c.setFont('LSB', 7); c.drawString(ax0 + 2, PH - ay0 + 2.5, a['id'])
        h_tot += r.height + 8
    c.setLineWidth(1)
    return h_tot


def qty_page(c, f):
    areas = floor_areas(f)
    header(c, f'{ordn(f)} Floor - M-5{f} SD R2: Scope and Quantities',
           'LD = linear supply diffuser in plenum box O.D. 48x4x26 (4 ft) w/ 1/2" lining  |  '
           'VD/CO on every branch  |  scope starts at the new tap, mains existing to remain')
    tw = [38, 300, 40, 112, 92, 46, 210]
    rows = [['Area', 'Location / scope', 'LD qty', 'CFM each\n(total)', 'Branches', 'VD/CO\ntagged',
             'Perimeter main in bubbled area\n(existing to remain, not priced)']]
    for a in areas:
        b = brsum([a])
        brt = '<br/>'.join(f'{n} x {s}, {l} in total' for s, (n, l) in b.items())
        scope = (f'<b>{esc(a["loc"])}</b><br/>F&amp;I {a["n"]} perimeter linear diffuser{"s" if a["n"] > 1 else ""} '
                 f'with plenum box{"es" if a["n"] > 1 else ""}; {a["n"]} branch tap{"s" if a["n"] > 1 else ""} off '
                 f'main with VD/CO, insulation, soffit coordination.' + (f'<br/><i>{esc(a["note"])}</i>' if a['note'] else ''))
        rows.append([Paragraph(f'<b>{a["id"]}</b>', st), Paragraph(scope, st), a['n'],
                     Paragraph(f'{" / ".join(str(v) for v in a["cfm"])}<br/>({sum(a["cfm"]):,})', str_),
                     Paragraph(brt, str_), a['vdco'] or '-', Paragraph(esc(a['main']), st)])
    b = brsum(areas)
    rows.append(['', f'{ordn(f)} floor total', sum(a['n'] for a in areas), f'{sum(sum(a["cfm"]) for a in areas):,}',
                 Paragraph('<br/>'.join(f'<b>{n} x {s}</b>' for s, (n, l) in b.items()), str_),
                 sum(a['vdco'] for a in areas), ''])
    h = table(c, rows, M, PH - TOP, tw)
    y = PH - TOP - h - 16
    y -= para(c, f'{ordn(f)} floor material list (for pricing)', sth, M, y, 600) + 4
    n = sum(a['n'] for a in areas); b = brsum(areas)
    rows = [['Item', 'Qty', 'Unit']]
    rows.append(['Perimeter linear supply diffuser, 48" (slots / finish per architect)', n, 'EA'])
    rows.append(['Plenum box O.D. 48x4x26 with 1/2" lining, side inlet', n, 'EA'])
    for s, (k, l) in b.items():
        rows.append([f'{s} branch duct, main tap to plenum box (as dimensioned, {l} in total)', k, 'EA'])
    rows.append(['Volume damper with operator (VD/CO) at each branch', sum(len(a['br']) for a in areas), 'EA'])
    rows.append(['Branch duct insulation, connections, hangers, sealing', sum(len(a['br']) for a in areas), 'EA'])
    rows.append(['Air balance of new diffusers', n, 'EA'])
    table(c, rows, M, y, [480, 60, 60], bold_last=False)
    x = M + sum(tw) + 16; w = PW - M - x; y = PH - TOP
    y -= para(c, 'Key plan', sth, x, y, w) + 3
    key_plans(c, f, x, y, w)
    footer(c); c.showPage()


def build_pdf():
    c = canvas.Canvas(OUT + '.pdf', pagesize=(PW, PH))
    c.setTitle('Perimeter Linear Diffusers - Breakdown by Floor'); c.setAuthor('Northern Wolves AC')
    summary_page(c)
    for f in FLOORS:
        screens_page(c, f)
        qty_page(c, f)
    c.save()


def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'Quantities by area'
    ws.append(['Snowflake 7 Times Square fl. 29-31 - Perimeter linear diffusers (SD R2 clouds) - Northern Wolves AC - ' + DATE])
    ws.append([])
    hdr = ['Floor', 'Area', 'Location', 'LD qty', 'CFM each', 'CFM total', '12x10 qty', '12x10 total in', '10x8 qty',
           '10x8 total in', 'VD/CO tagged', 'Main in bubbled area (ref.)', 'Notes',
           'Material $', 'Labor MH', 'Total $']
    ws.append(hdr)
    for cell in ws[3]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='1F3A5F')
        cell.alignment = Alignment(wrap_text=True, horizontal='center')
    row = 4; frows = []
    for f in FLOORS:
        start = row
        for a in floor_areas(f):
            b = brsum([a])
            ws.append([f, a['id'], a['loc'], a['n'], ' / '.join(map(str, a['cfm'])), sum(a['cfm']),
                       b.get('12x10', (0, 0))[0], b.get('12x10', (0, 0))[1], b.get('10x8', (0, 0))[0],
                       b.get('10x8', (0, 0))[1], a['vdco'], a['main'], a['note'], None, None, None])
            row += 1
        ws.append([f, '', f'{ordn(f)} floor total', f'=SUM(D{start}:D{row - 1})', '', f'=SUM(F{start}:F{row - 1})',
                   f'=SUM(G{start}:G{row - 1})', f'=SUM(H{start}:H{row - 1})', f'=SUM(I{start}:I{row - 1})',
                   f'=SUM(J{start}:J{row - 1})', f'=SUM(K{start}:K{row - 1})', '', '',
                   f'=SUM(N{start}:N{row - 1})', f'=SUM(O{start}:O{row - 1})', f'=SUM(P{start}:P{row - 1})'])
        for cell in ws[row]:
            cell.font = Font(bold=True)
        frows.append(row); row += 1
    ws.append(['', '', 'TOTAL'] + ['=' + '+'.join(f'{col}{r}' for r in frows) if col not in 'EL M' else ''
                                   for col in 'DEFGHIJKLMNOP'])
    for cell in ws[row]:
        cell.font = Font(bold=True)
    for col, w in zip('ABCDEFGHIJKLMNOP', [6, 6, 44, 7, 26, 9, 9, 11, 8, 11, 9, 48, 40, 11, 9, 11]):
        ws.column_dimensions[col].width = w
    for r in ws.iter_rows(min_row=4, max_row=row):
        for cell in r:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D4'
    wb.save(OUT + '.xlsx')


if __name__ == '__main__':
    build_pdf(); build_xlsx()
    print('ok', OUT)
