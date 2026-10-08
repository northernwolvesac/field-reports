"""Snowflake 7 Times Square fl. 29-31: perimeter linear diffusers - priced change order breakdown by floor.

Quantities and drawing crops come from build_breakdown.py (shop drawings M-529 / M-530 / M-531 SD R2 clouds).
Pricing basis per Ruslan 10/06/2026: linear diffuser $90/LF (48" = 4 LF), mechanic 5.5 MH per diffuser at $125,
foreman 30 / 15 / 15 hr at $155, air balancing $2,500 / $1,200 / $1,200, shop drawings $1,500 / $750 / $750.
Plenum box, branch and VD/CO unit costs are our estimate (not changed by Ruslan). No CO number yet.
Existing perimeter mains remain; work starts at the new taps. Deleted Bulletin 3 G diffusers: no note, no credit.

Run:  python3 build_co.py
Out:  Perimeter Linear Diffusers - Change Order Breakdown.pdf / .xlsx
"""
import importlib.util, os
from reportlab.lib import colors
from reportlab.platypus import Paragraph
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
import pymupdf

HERE = os.path.dirname(os.path.abspath(__file__))
spec = importlib.util.spec_from_file_location('bb', os.path.join(HERE, 'build_breakdown.py'))
bb = importlib.util.module_from_spec(spec); spec.loader.exec_module(bb)
OUT = os.path.join(HERE, 'Perimeter Linear Diffusers - Change Order Breakdown')
DATE = '10/06/2026'
PW, PH, M, TOP, BOT = bb.PW, bb.PH, bb.M, bb.TOP, bb.BOT
st, sth, esc, ordn = bb.st, bb.sth, bb.esc, bb.ordn

# ---- pricing ----
LD_LF, LD_LEN = 90, 4                  # $/LF, LF per diffuser
PLENUM = 95                            # plenum box O.D. 48x4x26 lined, shop fabricated
BRANCH = 45                            # 12x10 / 10x8 branch: tap, fittings, insulation
VDCO = 85                              # volume damper with concealed operator
MH_LD, RATE = 5.5, 125                 # mechanic MH per diffuser (install LD, plenum, tap, branch, VD/CO)
FOREMAN_RATE = 155
FOREMAN = {29: 30, 30: 15, 31: 15}     # hours
TAB = {29: 2500, 30: 1200, 31: 1200}
DRAFT = {29: 1500, 30: 750, 31: 750}
MAT_EA = LD_LF * LD_LEN + PLENUM + BRANCH + VDCO


def area_cost(a):
    return dict(mat=a['n'] * MAT_EA, mh=a['n'] * MH_LD, lab=a['n'] * MH_LD * RATE)


def floor_cost(f):
    ar = bb.floor_areas(f); n = sum(a['n'] for a in ar)
    d = dict(n=n, ld=n * LD_LF * LD_LEN, plenum=n * PLENUM, branch=n * BRANCH, vdco=n * VDCO,
             mh=n * MH_LD, lab=n * MH_LD * RATE, fh=FOREMAN[f], fore=FOREMAN[f] * FOREMAN_RATE,
             tab=TAB[f], draft=DRAFT[f])
    d['mat'] = d['ld'] + d['plenum'] + d['branch'] + d['vdco']
    d['total'] = d['mat'] + d['lab'] + d['fore'] + d['tab'] + d['draft']
    return d


FC = {f: floor_cost(f) for f in bb.FLOORS}
TOTAL = sum(d['total'] for d in FC.values())
assert [FC[f]['total'] for f in bb.FLOORS] == [37917.5, 19545, 22090] and TOTAL == 79552.5, (FC, TOTAL)
money = lambda v: f'${v:,.2f}' if v % 1 else f'${v:,.0f}'
NPAGES = 1 + 2 * len(bb.FLOORS)
page_no = [0]


def header(c, title, sub):
    bb.header(c, title, sub)


def footer(c):
    page_no[0] += 1
    c.setFillColor(colors.gray); c.setFont('LS', 8)
    c.drawString(M, 16, f'Northern Wolves AC  |  55 9th St, 55-A2, Brooklyn, NY 11215  |  (347) 463-9248  |  {DATE}')
    c.drawRightString(PW - M, 16, f'Page {page_no[0]} of {NPAGES}')


def summary_page(c):
    header(c, 'Change Order - Perimeter Linear Diffusers',
           f'Project 11010174 - Snowflake Fitout, 7 Times Square 29th-31st Floor  |  Submitted to Structure Tone  |  {DATE}')
    x = M; y = PH - TOP; w = 540
    y -= bb.para(c, 'Scope', sth, x, y, w) + 4
    notes = [
        'Furnish and install perimeter linear supply diffusers (48") along the window walls on floors 29, 30 and 31, '
        'as shown in the clouds on shop drawings M-529 / M-530 / M-531 SD R2.',
        'Each diffuser in a plenum box O.D. 48x4x26 with 1/2" lining, connected with a new 12x10 branch (one 10x8) '
        'with volume damper and concealed operator (VD/CO), tapped into the existing perimeter supply main.',
        'Existing perimeter mains remain; work starts at the new taps. Branch duct insulated, sealed and supported.',
        'Air balancing of the new diffusers and revised shop drawings included per floor.',
    ]
    for i, n in enumerate(notes, 1):
        y -= bb.para(c, f'{i}. {n}', st, x, y, w) + 5
    y -= 8
    y -= bb.para(c, 'Exclusions', sth, x, y, w) + 4
    for i, n in enumerate(['Soffit and ceiling openings, framing, access panels, patching and painting.',
                           'Relocation of any lights, electrical wires, j-boxes, or any other trades work that could be '
                           'discovered after soffits will be opened.'], 1):
        y -= bb.para(c, f'{i}. {n}', st, x, y, w) + 5

    x = 610; y = PH - TOP; w = PW - M - x
    y -= bb.para(c, 'Summary by floor', sth, x, y, w) + 4
    rows = [['Floor', 'Linear\ndiffusers', 'Material', 'Labor', 'Foreman', 'Air\nbalancing', 'Shop\ndrawings', 'Total']]
    for f in bb.FLOORS:
        d = FC[f]
        rows.append([f'{ordn(f)} floor (M-5{f})', d['n'], money(d['mat']), money(d['lab']), money(d['fore']),
                     money(d['tab']), money(d['draft']), money(d['total'])])
    k = lambda key: sum(FC[f][key] for f in bb.FLOORS)
    rows.append(['Total', k('n'), money(k('mat')), money(k('lab')), money(k('fore')), money(k('tab')),
                 money(k('draft')), money(TOTAL)])
    y -= bb.table(c, rows, x, y, [104, 52, 66, 72, 62, 62, 62, 74]) + 8
    y -= bb.para(c, f'Labor {k("mh"):g} MH at ${RATE}.00; foreman {k("fh")} hr at ${FOREMAN_RATE}.00.', st, x, y, w) + 18
    y -= bb.para(c, 'Change areas', sth, x, y, w) + 4
    rows = [['Area', 'Location', 'Diffusers', 'CFM']]
    for a in bb.AREAS:
        rows.append([a['id'], Paragraph(esc(a['loc']), st), a['n'], f'{sum(a["cfm"]):,}'])
    rows.append(['', 'Total', sum(a['n'] for a in bb.AREAS), f'{sum(sum(a["cfm"]) for a in bb.AREAS):,}'])
    bb.table(c, rows, x, y, [44, 366, 72, 72])
    footer(c); c.showPage()


def screens_page(c, f):
    areas = bb.floor_areas(f)
    header(c, f'{ordn(f)} Floor - M-5{f} SD R2: Change Areas',
           f'{FC[f]["n"]} perimeter linear diffusers in plenum boxes O.D. 48x4x26 w/ 1/2" lining, '
           f'{sum(sum(a["cfm"]) for a in areas):,} CFM  |  scope and cost on next page')
    for a, cx, cy, cw, ch in bb.layout(areas):
        c.setStrokeColor(bb.GRAY); c.setLineWidth(0.8); c.rect(cx, PH - cy - ch, cw, ch)
        c.setFillColor(bb.ORANGE); c.rect(cx, PH - cy - bb.TB, cw, bb.TB, stroke=0, fill=1)
        c.setFillColor(colors.white); c.setFont('LSB', 9)
        c.drawString(cx + 5, PH - cy - 11.5, f'{a["id"]}  {a["loc"]}' if cw > 330 else a['id'])
        c.drawRightString(cx + cw - 5, PH - cy - 11.5, f'{a["n"]} x LD  |  {sum(a["cfm"]):,} CFM')
        cr = bb.crop_rect(a); p = bb.DOCS[f]['AB'.index(a['sheet'])]
        r = bb.fit(cr, pymupdf.Rect(cx + 3, cy + bb.TB + 4, cx + cw - 3, cy + ch - 4))
        dpi = int(min(300, max(110, 72 * r.width / cr.width * 2.2)))
        c.drawImage(ImageReader(bb.png(p, cr, dpi)), r.x0, PH - r.y1, r.width, r.height)
        c.setFont('LS', 7.5); c.setFillColor(bb.GRAY)
        c.drawRightString(cx + cw - 5, PH - cy - ch + 4, f'M-5{f}{a["sheet"]}.SD')
    footer(c); c.showPage()


def cost_page(c, f):
    areas = bb.floor_areas(f); d = FC[f]
    header(c, f'{ordn(f)} Floor - M-5{f} SD R2: Scope of Work and Cost Breakdown',
           f'Labor at ${RATE}.00 per man-hour, foreman ${FOREMAN_RATE}.00 per hour  |  '
           'existing perimeter mains remain, work starts at the new taps')
    tw = [38, 352, 40, 96, 70, 76, 70]
    rows = [['Area', 'Scope of work', 'LD qty', 'CFM each\n(total)', 'Material', 'Labor', 'Total']]
    for a in areas:
        ac = area_cost(a)
        br = bb.brsum([a])
        brt = ', '.join(f'{n} x {s}' for s, (n, l) in br.items())
        scope = (f'<b>{esc(a["loc"])}</b><br/>F&amp;I {a["n"]} perimeter linear diffuser{"s" if a["n"] > 1 else ""} '
                 f'48" with plenum box{"es" if a["n"] > 1 else ""} O.D. 48x4x26 lined; new tap{"s" if a["n"] > 1 else ""} '
                 f'on existing main, {brt} branch with VD/CO, insulation.')
        rows.append([Paragraph(f'<b>{a["id"]}</b>', st), Paragraph(scope, st), a['n'],
                     Paragraph(f'{" / ".join(str(v) for v in a["cfm"])}<br/>({sum(a["cfm"]):,})', bb.str_),
                     money(ac['mat']), Paragraph(f'{ac["mh"]:g} MH<br/>{money(ac["lab"])}', bb.str_),
                     money(ac['mat'] + ac['lab'])])
    rows.append(['', 'Foreman (layout, soffit coordination, supervision)', '', '', '',
                 Paragraph(f'{d["fh"]} hr<br/>{money(d["fore"])}', bb.str_), money(d['fore'])])
    rows.append(['', 'Air balancing of new diffusers', '', '', '', '', money(d['tab'])])
    rows.append(['', 'Shop drawings for revised design (SD R2)', '', '', '', '', money(d['draft'])])
    rows.append(['', f'{ordn(f)} floor total', d['n'], f'{sum(sum(a["cfm"]) for a in areas):,}', money(d['mat']),
                 money(d['lab'] + d['fore']), money(d['total'])])
    h = bb.table(c, rows, M, PH - TOP, tw, extra=[('FONT', (-1, 1), (-1, -1), 'LSB', 9),
                                                  ('VALIGN', (0, -4), (-1, -1), 'MIDDLE')])
    y = PH - TOP - h - 16
    y -= bb.para(c, f'{ordn(f)} floor cost build-up', sth, M, y, 600) + 4
    rows = [['Item', 'Qty', 'Unit', 'Rate', 'Amount']]
    rows.append([f'Perimeter linear supply diffuser 48" ({LD_LEN} LF each)', d['n'] * LD_LEN, 'LF', money(LD_LF), money(d['ld'])])
    rows.append(['Plenum box O.D. 48x4x26 with 1/2" lining', d['n'], 'EA', money(PLENUM), money(d['plenum'])])
    rows.append(['Branch duct from new tap to plenum box, fittings, insulation', d['n'], 'EA', money(BRANCH), money(d['branch'])])
    rows.append(['Volume damper with concealed operator (VD/CO)', d['n'], 'EA', money(VDCO), money(d['vdco'])])
    rows.append(['Labor: install diffusers, plenum boxes, taps, branches, VD/CO', f'{d["mh"]:g}', 'MH', money(RATE), money(d['lab'])])
    rows.append(['Foreman', d['fh'], 'HR', money(FOREMAN_RATE), money(d['fore'])])
    rows.append(['Air balancing', 1, 'LS', '', money(d['tab'])])
    rows.append(['Shop drawings', 1, 'LS', '', money(d['draft'])])
    rows.append([f'{ordn(f)} floor total', '', '', '', money(d['total'])])
    bb.table(c, rows, M, y, [420, 50, 40, 70, 90])
    x = M + sum(tw) + 16; w = PW - M - x; y = PH - TOP
    y -= bb.para(c, 'Key plan', sth, x, y, w) + 3
    y -= bb.key_plans(c, f, x, y, w) + 6
    rows = [['Floor', 'Total']] + [[f'{ordn(g)} floor', money(FC[g]['total'])] for g in bb.FLOORS] + [['Total', money(TOTAL)]]
    i = bb.FLOORS.index(f) + 1
    bb.table(c, rows, x, y, [w - 90, 90], extra=[('BACKGROUND', (0, i), (-1, i), colors.HexColor('#FCEBD2'))])
    footer(c); c.showPage()


def build_pdf():
    c = canvas.Canvas(OUT + '.pdf', pagesize=(PW, PH))
    c.setTitle('Change Order - Perimeter Linear Diffusers'); c.setAuthor('Northern Wolves AC')
    summary_page(c)
    for f in bb.FLOORS:
        screens_page(c, f)
        cost_page(c, f)
    c.save()


def build_xlsx():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = 'Rates'
    hdr_fill = PatternFill('solid', fgColor='1F3A5F')
    rates = [('Linear diffuser $/LF', LD_LF), ('LF per diffuser', LD_LEN), ('Plenum box $/EA', PLENUM),
             ('Branch $/EA', BRANCH), ('VD/CO $/EA', VDCO), ('Mechanic MH per diffuser', MH_LD),
             ('Mechanic $/MH', RATE), ('Foreman $/HR', FOREMAN_RATE)]
    ws.append(['Rate', 'Value'])
    for cell in ws[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = hdr_fill
    for r in rates:
        ws.append(list(r))
    ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 10
    R = {k: f'Rates!$B${i + 2}' for i, (k, _) in enumerate(rates)}
    mat = f"({R['Linear diffuser $/LF']}*{R['LF per diffuser']}+{R['Plenum box $/EA']}+{R['Branch $/EA']}+{R['VD/CO $/EA']})"

    ws2 = wb.create_sheet('By area'); ws2.append(['Floor', 'Area', 'Location', 'LD qty', 'CFM', 'Material $', 'Labor MH', 'Labor $', 'Total $'])
    for cell in ws2[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = hdr_fill
    row = 2
    for a in bb.AREAS:
        ws2.append([a['floor'], a['id'], a['loc'], a['n'], sum(a['cfm']), f'=D{row}*{mat}',
                    f"=D{row}*{R['Mechanic MH per diffuser']}", f"=G{row}*{R['Mechanic $/MH']}", f'=F{row}+H{row}'])
        row += 1
    for col, w in zip('ABCDEFGHI', [6, 6, 48, 7, 8, 11, 9, 11, 11]):
        ws2.column_dimensions[col].width = w

    ws3 = wb.create_sheet('By floor')
    ws3.append(['Floor', 'LD qty', 'Material $', 'Labor $', 'Foreman hr', 'Foreman $', 'Air balancing $', 'Shop drawings $', 'Total $'])
    for cell in ws3[1]:
        cell.font = Font(bold=True, color='FFFFFF'); cell.fill = hdr_fill
    for i, f in enumerate(bb.FLOORS, 2):
        rng = lambda col: f"SUMIF('By area'!A:A,A{i},'By area'!{col}:{col})"
        ws3.append([f, f'={rng("D")}', f'={rng("F")}', f'={rng("H")}', FOREMAN[f], f"=E{i}*{R['Foreman $/HR']}",
                    TAB[f], DRAFT[f], f'=C{i}+D{i}+F{i}+G{i}+H{i}'])
    n = len(bb.FLOORS) + 1
    ws3.append(['Total'] + [f'=SUM({col}2:{col}{n})' for col in 'BCDEFGHI'])
    for cell in ws3[n + 1]:
        cell.font = Font(bold=True)
    for ws_ in (ws2, ws3):
        for r in ws_.iter_rows(min_row=2):
            for cell in r:
                if cell.column_letter in ('CDFGHI' if ws_ is ws3 else 'FHI') and not (ws_ is ws3 and cell.column_letter == 'E'):
                    cell.number_format = '$#,##0.00'
    ws3.column_dimensions['A'].width = 8
    for col in 'BCDEFGHI':
        ws3.column_dimensions[col].width = 14
    wb.save(OUT + '.xlsx')


if __name__ == '__main__':
    build_pdf(); build_xlsx()
    for f in bb.FLOORS:
        print(f, {k: v for k, v in FC[f].items()})
    print('TOTAL', TOTAL, 'ok', OUT)
